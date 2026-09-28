from csv import writer
from dataclasses import dataclass
from datetime import (
    UTC,
    datetime,
)
from io import (
    BytesIO,
    StringIO,
)
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import (
    A4,
    landscape,
)
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


@dataclass(frozen=True)
class ArquivoExportado:
    conteudo: bytes
    media_type: str
    nome_arquivo: str


class ExportacaoRelatoriosService:
    """Transforma dados de relatórios em arquivos para download."""

    CABECALHOS = (
        "ID",
        "Tipo",
        "Modelo",
        "Aluguéis",
        "Receita (R$)",
        "Manutenções",
        "Custos (R$)",
        "Resultado bruto (R$)",
    )

    def __init__(
        self,
        relatorio_service,
    ):
        if relatorio_service is None:
            raise ValueError(
                "RelatorioService é obrigatório."
            )

        self.relatorio_service = (
            relatorio_service
        )

    @staticmethod
    def _texto_planilha(
        valor,
    ):
        texto = str(
            valor
            if valor is not None
            else ""
        )

        if (
            texto.lstrip()
            .startswith(
                ("=", "+", "-", "@")
            )
        ):
            return f"'{texto}"

        return texto

    @staticmethod
    def _texto_pdf(
        valor,
    ):
        texto = str(
            valor
            if valor is not None
            else ""
        )

        return (
            texto.encode(
                "cp1252",
                errors="replace",
            )
            .decode("cp1252")
        )

    @staticmethod
    def _moeda_brl(
        valor,
    ):
        numero = float(
            valor or 0
        )
        texto = (
            f"{numero:,.2f}"
            .replace(",", "_")
            .replace(".", ",")
            .replace("_", ".")
        )
        return f"R$ {texto}"

    @staticmethod
    def _instante_geracao():
        return datetime.now(
            UTC
        )

    def _nome_arquivo(
        self,
        formato,
        instante,
    ):
        sufixo = instante.strftime(
            "%Y%m%d_%H%M%S"
        )

        return (
            "locadora_resultado_por_veiculo_"
            f"{sufixo}.{formato}"
        )

    def _linhas(
        self,
    ):
        dados = (
            self.relatorio_service
            .resultado_por_veiculo()
        )

        return [
            (
                int(item["id"]),
                str(
                    item["tipo"]
                ),
                str(
                    item["modelo"]
                ),
                int(
                    item["total_alugueis"]
                ),
                float(
                    item["receita"]
                ),
                int(
                    item[
                        "total_manutencoes"
                    ]
                ),
                float(
                    item[
                        "custo_manutencao"
                    ]
                ),
                float(
                    item[
                        "resultado_bruto"
                    ]
                ),
            )
            for item in dados
        ]

    def exportar_resultado_por_veiculo(
        self,
        formato,
    ):
        formato = str(
            formato
        ).strip().lower()

        instante = (
            self._instante_geracao()
        )
        linhas = self._linhas()

        geradores = {
            "csv": (
                self._gerar_csv
            ),
            "xlsx": (
                self._gerar_xlsx
            ),
            "pdf": (
                self._gerar_pdf
            ),
        }

        if formato not in geradores:
            raise ValueError(
                "Formato de exportação inválido."
            )

        conteudo, media_type = (
            geradores[formato](
                linhas,
                instante,
            )
        )

        return ArquivoExportado(
            conteudo=conteudo,
            media_type=media_type,
            nome_arquivo=(
                self._nome_arquivo(
                    formato,
                    instante,
                )
            ),
        )

    def _gerar_csv(
        self,
        linhas,
        instante,
    ):
        arquivo = StringIO(
            newline=""
        )

        csv_writer = writer(
            arquivo,
            delimiter=";",
            lineterminator="\n",
        )

        csv_writer.writerow(
            self.CABECALHOS
        )

        for linha in linhas:
            csv_writer.writerow(
                (
                    linha[0],
                    self._texto_planilha(
                        linha[1]
                    ),
                    self._texto_planilha(
                        linha[2]
                    ),
                    *linha[3:],
                )
            )

        # BOM melhora a abertura direta em planilhas no Windows
        # sem perder a codificação UTF-8 dos textos em português.
        conteudo = (
            "\ufeff"
            + arquivo.getvalue()
        ).encode("utf-8")

        return (
            conteudo,
            "text/csv",
        )

    def _gerar_xlsx(
        self,
        linhas,
        instante,
    ):
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = (
            "Resultado por veículo"
        )

        planilha.merge_cells(
            "A1:H1"
        )
        planilha["A1"] = (
            "Locadora — Resultado por veículo"
        )
        planilha["A1"].font = Font(
            bold=True,
            size=16,
            color="FFFFFF",
        )
        planilha["A1"].fill = (
            PatternFill(
                fill_type="solid",
                fgColor="0D2033",
            )
        )
        planilha["A1"].alignment = (
            Alignment(
                horizontal="center"
            )
        )

        planilha.merge_cells(
            "A2:H2"
        )
        planilha["A2"] = (
            "Gerado em UTC: "
            + instante.strftime(
                "%d/%m/%Y %H:%M:%S"
            )
        )
        planilha["A2"].font = Font(
            italic=True,
            color="526579",
        )

        linha_cabecalho = 4

        for coluna, cabecalho in enumerate(
            self.CABECALHOS,
            start=1,
        ):
            celula = planilha.cell(
                row=linha_cabecalho,
                column=coluna,
                value=cabecalho,
            )
            celula.font = Font(
                bold=True,
                color="FFFFFF",
            )
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="0F8B83",
            )
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for indice, linha in enumerate(
            linhas,
            start=linha_cabecalho + 1,
        ):
            for coluna, valor in enumerate(
                linha,
                start=1,
            ):
                if coluna in (2, 3):
                    valor = (
                        self._texto_planilha(
                            valor
                        )
                    )

                planilha.cell(
                    row=indice,
                    column=coluna,
                    value=valor,
                )

            for coluna in (
                5,
                7,
                8,
            ):
                planilha.cell(
                    row=indice,
                    column=coluna,
                ).number_format = (
                    'R$ #,##0.00'
                )

        ultima_linha = max(
            linha_cabecalho,
            linha_cabecalho
            + len(linhas),
        )

        planilha.freeze_panes = (
            "A5"
        )
        planilha.auto_filter.ref = (
            f"A4:H{ultima_linha}"
        )

        larguras = (
            10,
            18,
            28,
            14,
            18,
            16,
            18,
            22,
        )

        for indice, largura in enumerate(
            larguras,
            start=1,
        ):
            planilha.column_dimensions[
                chr(64 + indice)
            ].width = largura

        planilha.row_dimensions[1].height = 28
        planilha.row_dimensions[4].height = 24

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        return (
            buffer.getvalue(),
            (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def _gerar_pdf(
        self,
        linhas,
        instante,
    ):
        buffer = BytesIO()

        documento = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title=(
                "Locadora - Resultado por veículo"
            ),
            author="Locadora",
        )

        estilos = (
            getSampleStyleSheet()
        )
        titulo = ParagraphStyle(
            "TituloLocadora",
            parent=estilos["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor(
                "#0D2033"
            ),
            alignment=TA_CENTER,
            spaceAfter=5 * mm,
        )
        subtitulo = ParagraphStyle(
            "SubtituloLocadora",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=8,
            textColor=colors.HexColor(
                "#526579"
            ),
            alignment=TA_CENTER,
            spaceAfter=5 * mm,
        )
        texto_tabela = ParagraphStyle(
            "TextoTabela",
            parent=estilos["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
        )

        elementos = [
            Paragraph(
                "Locadora — Resultado por veículo",
                titulo,
            ),
            Paragraph(
                (
                    "Gerado em UTC: "
                    + instante.strftime(
                        "%d/%m/%Y %H:%M:%S"
                    )
                ),
                subtitulo,
            ),
        ]

        dados_tabela = [
            [
                Paragraph(
                    escape(cabecalho),
                    texto_tabela,
                )
                for cabecalho
                in self.CABECALHOS
            ]
        ]

        for linha in linhas:
            dados_tabela.append(
                [
                    str(linha[0]),
                    Paragraph(
                        escape(
                            self._texto_pdf(
                                linha[1]
                            )
                        ),
                        texto_tabela,
                    ),
                    Paragraph(
                        escape(
                            self._texto_pdf(
                                linha[2]
                            )
                        ),
                        texto_tabela,
                    ),
                    str(linha[3]),
                    self._moeda_brl(
                        linha[4]
                    ),
                    str(linha[5]),
                    self._moeda_brl(
                        linha[6]
                    ),
                    self._moeda_brl(
                        linha[7]
                    ),
                ]
            )

        if not linhas:
            dados_tabela.append(
                [
                    Paragraph(
                        (
                            "Nenhum resultado disponível "
                            "para exportação."
                        ),
                        texto_tabela,
                    ),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )

        tabela = Table(
            dados_tabela,
            repeatRows=1,
            colWidths=[
                14 * mm,
                25 * mm,
                42 * mm,
                20 * mm,
                29 * mm,
                24 * mm,
                29 * mm,
                33 * mm,
            ],
        )

        estilos_tabela = [
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor(
                    "#0F8B83"
                ),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "RIGHT",
            ),
            (
                "ALIGN",
                (1, 1),
                (2, -1),
                "LEFT",
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.25,
                colors.HexColor(
                    "#CBD5DF"
                ),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor(
                        "#F7FAFC"
                    ),
                ],
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ]

        if not linhas:
            estilos_tabela.extend(
                [
                    (
                        "SPAN",
                        (0, 1),
                        (-1, 1),
                    ),
                    (
                        "ALIGN",
                        (0, 1),
                        (-1, 1),
                        "CENTER",
                    ),
                ]
            )

        tabela.setStyle(
            TableStyle(
                estilos_tabela
            )
        )

        elementos.extend(
            [
                tabela,
                Spacer(1, 2 * mm),
            ]
        )

        documento.build(
            elementos
        )

        return (
            buffer.getvalue(),
            "application/pdf",
        )
