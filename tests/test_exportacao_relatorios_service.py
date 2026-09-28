from io import BytesIO

import pytest

from openpyxl import load_workbook

from modulos.servicos.exportacao_relatorios_service import (
    ExportacaoRelatoriosService,
)


class RelatorioServiceFake:
    def __init__(
        self,
        dados=None,
    ):
        self.dados = (
            dados
            if dados is not None
            else [
                {
                    "id": 1,
                    "tipo": "carro",
                    "modelo": "Corolla",
                    "total_alugueis": 4,
                    "receita": 3200.0,
                    "total_manutencoes": 2,
                    "custo_manutencao": 700.0,
                    "resultado_bruto": 2500.0,
                },
                {
                    "id": 2,
                    "tipo": "moto",
                    "modelo": "CB 500",
                    "total_alugueis": 3,
                    "receita": 1800.0,
                    "total_manutencoes": 1,
                    "custo_manutencao": 300.0,
                    "resultado_bruto": 1500.0,
                },
            ]
        )

    def resultado_por_veiculo(
        self,
    ):
        return list(
            self.dados
        )


def criar_service(
    dados=None,
):
    return ExportacaoRelatoriosService(
        relatorio_service=(
            RelatorioServiceFake(
                dados
            )
        )
    )


def test_exige_relatorio_service():
    with pytest.raises(
        ValueError,
        match="RelatorioService",
    ):
        ExportacaoRelatoriosService(
            relatorio_service=None
        )


def test_exporta_csv_utf8_com_cabecalho():
    arquivo = (
        criar_service()
        .exportar_resultado_por_veiculo(
            "csv"
        )
    )

    texto = arquivo.conteudo.decode(
        "utf-8-sig"
    )

    assert arquivo.media_type == "text/csv"
    assert arquivo.nome_arquivo.endswith(
        ".csv"
    )
    assert (
        "ID;Tipo;Modelo;Aluguéis;"
        "Receita (R$);Manutenções;"
        "Custos (R$);Resultado bruto (R$)"
        in texto
    )
    assert (
        "1;carro;Corolla;4;3200.0;2;700.0;2500.0"
        in texto
    )


def test_csv_e_xlsx_neutralizam_formula_injetada():
    dados = [
        {
            "id": 1,
            "tipo": "=HYPERLINK",
            "modelo": "+SUM(1,1)",
            "total_alugueis": 1,
            "receita": 100.0,
            "total_manutencoes": 0,
            "custo_manutencao": 0.0,
            "resultado_bruto": 100.0,
        }
    ]

    service = criar_service(
        dados
    )

    csv_arquivo = (
        service
        .exportar_resultado_por_veiculo(
            "csv"
        )
    )
    csv_texto = (
        csv_arquivo.conteudo
        .decode("utf-8-sig")
    )

    assert "'=HYPERLINK" in csv_texto
    assert "'+SUM(1,1)" in csv_texto

    xlsx_arquivo = (
        service
        .exportar_resultado_por_veiculo(
            "xlsx"
        )
    )
    workbook = load_workbook(
        BytesIO(
            xlsx_arquivo.conteudo
        )
    )
    planilha = workbook[
        "Resultado por veículo"
    ]

    assert planilha["B5"].value == "'=HYPERLINK"
    assert planilha["C5"].value == "'+SUM(1,1)"
    assert planilha["B5"].data_type != "f"
    assert planilha["C5"].data_type != "f"

    workbook.close()


def test_exporta_xlsx_valido_e_formatado():
    arquivo = (
        criar_service()
        .exportar_resultado_por_veiculo(
            "xlsx"
        )
    )

    assert arquivo.media_type == (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )
    assert arquivo.nome_arquivo.endswith(
        ".xlsx"
    )

    workbook = load_workbook(
        BytesIO(
            arquivo.conteudo
        )
    )
    planilha = workbook[
        "Resultado por veículo"
    ]

    assert planilha["A1"].value == (
        "Locadora — Resultado por veículo"
    )
    assert planilha["A4"].value == "ID"
    assert planilha["C5"].value == "Corolla"
    assert planilha["E5"].value == 3200.0
    assert planilha["E5"].number_format == (
        "R$ #,##0.00"
    )
    assert planilha.freeze_panes == "A5"
    assert planilha.auto_filter.ref == "A4:H6"

    workbook.close()


def test_exporta_pdf_valido():
    arquivo = (
        criar_service()
        .exportar_resultado_por_veiculo(
            "pdf"
        )
    )

    assert arquivo.media_type == "application/pdf"
    assert arquivo.nome_arquivo.endswith(
        ".pdf"
    )
    assert arquivo.conteudo.startswith(
        b"%PDF-"
    )
    assert arquivo.conteudo.rstrip().endswith(
        b"%%EOF"
    )
    assert len(arquivo.conteudo) > 1000


def test_exporta_arquivos_vazios_sem_falhar():
    service = criar_service(
        []
    )

    csv_arquivo = (
        service
        .exportar_resultado_por_veiculo(
            "csv"
        )
    )
    xlsx_arquivo = (
        service
        .exportar_resultado_por_veiculo(
            "xlsx"
        )
    )
    pdf_arquivo = (
        service
        .exportar_resultado_por_veiculo(
            "pdf"
        )
    )

    assert len(
        csv_arquivo.conteudo
    ) > 10
    assert len(
        xlsx_arquivo.conteudo
    ) > 1000
    assert pdf_arquivo.conteudo.startswith(
        b"%PDF-"
    )


def test_formato_invalido_falha():
    with pytest.raises(
        ValueError,
        match="Formato",
    ):
        (
            criar_service()
            .exportar_resultado_por_veiculo(
                "xml"
            )
        )
